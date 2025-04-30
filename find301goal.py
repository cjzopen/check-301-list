import openpyxl
import requests

def get_final_url(url):
  """
  追蹤 301 重定向並返回最終的目的地網址。
  """
  try:
    response = requests.head(url, allow_redirects=True, timeout=20)  # Increased timeout to 20 seconds
    return response.url
  except requests.RequestException as e:
    print(f"Error processing URL {url}: {e}. Please check the URL or your network connection.")
    return None

def process_excel(file_path):
  """
  處理 Excel 檔案，將 B 欄的網址追蹤 301 後的目的地網址，並填入 C 欄。
  """
  # 開啟 Excel 檔案
  workbook = openpyxl.load_workbook(file_path)
  sheet = workbook.active

  # 從第二行開始處理資料
  for row in sheet.iter_rows(min_row=2, max_row=sheet.max_row, min_col=2, max_col=2):
    cell = row[0]
    if cell.value:  # 確保 B 欄有值
      final_url = get_final_url(cell.value)
    if final_url:
      # 將最終網址寫入 C 欄
      sheet.cell(row=cell.row, column=3, value=final_url)
      print(f"Row {cell.row}: Processed URL {cell.value} -> {final_url}")

  # 儲存修改後的 Excel 檔案
  workbook.save(file_path)
  print(f"Excel 檔案已更新: {file_path}")

if __name__ == "__main__":
  excel_file = "301.xlsx"
  process_excel(excel_file)